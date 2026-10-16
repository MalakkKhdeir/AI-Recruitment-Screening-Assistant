
from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_path):
    reader = PdfReader(str(file_path))

    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ValueError("PDF is password-protected.")
        except Exception as exc:
            raise ValueError(
                "Cannot read encrypted PDF."
            ) from exc

    pages_text = []

    for page in reader.pages:
        pages_text.append(page.extract_text() or "")

    return "\n".join(pages_text).strip()


def extract_text_from_docx(file_path):
    document = Document(str(file_path))
    parts = []

    # Extract text from paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    # Extract text from tables (common in CVs)
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_text(file_path):
    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        text = extract_text_from_pdf(file_path)

    elif suffix == ".docx":
        text = extract_text_from_docx(file_path)

    else:
        raise ValueError(
            f"Unsupported file format: {suffix}"
        )

    # Normalize whitespace while preserving line breaks
    lines = [
        " ".join(line.split())
        for line in text.splitlines()
        if line.strip()
    ]

    return "\n".join(lines)


if __name__ == "__main__":
    attachments_dir = Path("attachments")

    files = sorted(
        list(attachments_dir.glob("*.pdf"))
        + list(attachments_dir.glob("*.docx"))
        + list(attachments_dir.glob("*.PDF"))
        + list(attachments_dir.glob("*.DOCX"))
    )

    if not files:
        print("No CV files found.")

    for file in files:
        print(f"\nProcessing: {file.name}")
        print("-" * 40)

        try:
            text = extract_text(file)

            if text:
                print(f"Extracted characters: {len(text)}")
                print(text)
            else:
                print("No text extracted from this file.")

        except Exception as exc:
            print(f"Error processing {file.name}: {exc}")
            continue
